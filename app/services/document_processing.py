import logging
from pathlib import Path
from time import perf_counter

from docx import Document as WordDocument
from pypdf import PdfReader
from sqlalchemy.orm import Session, sessionmaker

from app.db.session import get_session_factory
from app.models.document import Document
from app.models.document_chunk import DocumentChunk
from app.models.document_processing_task import DocumentProcessingTask
from app.core.timezone import beijing_now
from app.services.file_storage import resolve_storage_path

CHUNK_SIZE = 1000
CHUNK_OVERLAP = 150
MAX_CHUNKS_PER_DOCUMENT = 2000
logger = logging.getLogger("repomind.document_processing")


def extract_text(file_path: Path) -> str:
    suffix = file_path.suffix.lower()
    if suffix in {".md", ".markdown", ".txt", ".json", ".yaml", ".yml"}:
        return read_text_file(file_path)
    if suffix == ".pdf":
        return extract_pdf_text(file_path)
    if suffix == ".docx":
        return extract_docx_text(file_path)
    raise ValueError("暂不支持该文件类型")


def read_text_file(file_path: Path) -> str:
    for encoding in ("utf-8-sig", "utf-8", "gb18030"):
        try:
            return file_path.read_text(encoding=encoding)
        except UnicodeDecodeError:
            continue
    raise ValueError("无法识别文本文件编码")


def extract_pdf_text(file_path: Path) -> str:
    reader = PdfReader(str(file_path))
    if reader.is_encrypted:
        raise ValueError("暂不支持加密 PDF")
    return "\n\n".join(page.extract_text() or "" for page in reader.pages)


def extract_docx_text(file_path: Path) -> str:
    document = WordDocument(str(file_path))
    parts = [paragraph.text for paragraph in document.paragraphs if paragraph.text.strip()]
    for table in document.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells]
            if any(cells):
                parts.append(" | ".join(cells))
    return "\n\n".join(parts)


def split_text(text: str, chunk_size: int = CHUNK_SIZE, overlap: int = CHUNK_OVERLAP) -> list[tuple[str, int, int]]:
    normalized = text.replace("\r\n", "\n").replace("\r", "\n").strip()
    if not normalized:
        raise ValueError("文件中没有可处理的文本")

    chunks: list[tuple[str, int, int]] = []
    start = 0
    text_length = len(normalized)
    while start < text_length:
        end = min(start + chunk_size, text_length)
        if end < text_length:
            search_start = start + (chunk_size // 2)
            candidate_breaks = [
                normalized.rfind("\n", search_start, end),
                normalized.rfind("。", search_start, end),
                normalized.rfind(" ", search_start, end),
            ]
            best_break = max(candidate_breaks)
            if best_break > start:
                end = best_break + 1

        content = normalized[start:end].strip()
        if content:
            chunks.append((content, start, end))
        if len(chunks) > MAX_CHUNKS_PER_DOCUMENT:
            raise ValueError("文档分段数量超过上限，请缩小文件内容")
        if end >= text_length:
            break
        start = max(end - overlap, start + 1)
    return chunks


def process_document_task(
    task_id: int,
    session_factory: sessionmaker[Session] | None = None,
) -> None:
    db = (session_factory or get_session_factory())()
    started_at = perf_counter()
    try:
        task = db.get(DocumentProcessingTask, task_id)
        if task is None:
            return
        document = db.get(Document, task.document_id)
        if document is None or not document.storage_path:
            mark_task_failed(db, task_id, "找不到待处理的文档文件")
            return

        task.status = "running"
        task.started_at = beijing_now()
        document.processing_status = "processing"
        document.processing_error = None
        db.commit()

        try:
            text = extract_text(resolve_storage_path(document.storage_path))
            chunks = split_text(text)
        except (OSError, ValueError) as exc:
            logger.warning(
                "document_processing_failed",
                extra={
                    "task_id": task_id,
                    "document_id": document.id,
                    "operation": task.operation,
                    "latency_ms": round((perf_counter() - started_at) * 1000),
                    "error_type": type(exc).__name__,
                },
            )
            mark_task_failed(db, task_id, str(exc))
            return

        db.query(DocumentChunk).filter(DocumentChunk.document_id == document.id).delete()
        db.add_all(
            [
                DocumentChunk(
                    document_id=document.id,
                    chunk_index=index,
                    content=content,
                    char_start=char_start,
                    char_end=char_end,
                )
                for index, (content, char_start, char_end) in enumerate(chunks)
            ]
        )
        task.status = "completed"
        task.completed_at = beijing_now()
        task.error_message = None
        document.processing_status = "completed"
        document.processing_error = None
        document.processed_at = beijing_now()
        db.commit()
        logger.info(
            "document_processing_completed",
            extra={
                "task_id": task_id,
                "document_id": document.id,
                "operation": task.operation,
                "chunk_count": len(chunks),
                "latency_ms": round((perf_counter() - started_at) * 1000),
            },
        )
    except Exception as exc:
        db.rollback()
        logger.warning(
            "document_processing_failed",
            extra={
                "task_id": task_id,
                "latency_ms": round((perf_counter() - started_at) * 1000),
                "error_type": type(exc).__name__,
            },
        )
        mark_task_failed(db, task_id, "文本提取失败，请检查文件格式或内容")
    finally:
        db.close()


def mark_task_failed(db: Session, task_id: int, error_message: str) -> None:
    db.rollback()
    task = db.get(DocumentProcessingTask, task_id)
    if task is None:
        return
    document = db.get(Document, task.document_id)
    task.status = "failed"
    task.completed_at = beijing_now()
    task.error_message = error_message[:1000]
    if document is not None:
        document.processing_status = "failed"
        document.processing_error = error_message[:1000]
    db.commit()
